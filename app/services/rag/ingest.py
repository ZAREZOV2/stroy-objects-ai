"""Загрузка документов и ответы на вопросы."""

import io
import re
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.config import get_settings
from app.models.document import Document, DocumentChunk
from app.schemas.rag import AskResponse, Citation
from app.services import objects as object_service
from app.services.rag import chunking
from app.services.rag import ollama as ollama_service
from app.services.rag import search as search_service

settings = get_settings()

SYSTEM_PROMPT = (
    "Ты помощник сотрудника министерства строительства. "
    "Отвечай кратко и по делу, только на основе предоставленного контекста."
)


def extract_text(content: bytes, file_type: str) -> str:
    if file_type == "pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if file_type == "docx":
        doc = DocxDocument(io.BytesIO(content))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return content.decode("utf-8", errors="replace")


async def ingest_bytes(
    db: AsyncSession,
    *,
    title: str,
    filename: str,
    content: bytes,
    is_demo: bool = False,
) -> Document:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in ("pdf", "docx", "txt"):
        raise ValueError("Поддерживаются форматы PDF, DOCX, TXT")
    text = extract_text(content, ext)
    chunks_text = chunking.split_text(text)
    doc = Document(title=title, filename=filename, file_type=ext, is_demo=is_demo)
    db.add(doc)
    await db.flush()

    embeddings = await ollama_service.embed_texts(chunks_text)
    for i, chunk in enumerate(chunks_text):
        emb = embeddings[i] if embeddings else None
        db.add(
            DocumentChunk(
                document_id=doc.id,
                chunk_index=i,
                content=chunk,
                embedding=emb,
            )
        )
    await db.commit()
    await db.refresh(doc)
    return doc


async def list_documents(db: AsyncSession) -> list[tuple[Document, int]]:
    q = (
        select(Document, func.count(DocumentChunk.id))
        .outerjoin(DocumentChunk, DocumentChunk.document_id == Document.id)
        .group_by(Document.id)
        .order_by(Document.id.desc())
    )
    result = await db.execute(q)
    return [(row[0], row[1]) for row in result.all()]


async def get_document(db: AsyncSession, doc_id: int) -> Document | None:
    q = select(Document).where(Document.id == doc_id).options(selectinload(Document.chunks))
    return (await db.execute(q)).scalar_one_or_none()


def is_registry_objects_question(question: str) -> bool:
    """True только для явных вопросов к реестру (просрочка, список), не к нормативным документам."""
    q = question.strip()
    if not q:
        return False
    lower = q.lower()
    if re.search(r"просроч", lower):
        return True
    if re.search(r"список\s+.*объект", lower) or re.search(r"объект.*\s+в\s+реестре", lower):
        return True
    if "реестр" in lower and re.search(r"сколько|какие|список", lower):
        return True
    if re.search(r"какие\s+объект", lower) and re.search(r"просроч|реестр|строительств", lower):
        return True
    return False


async def answer_question(db: AsyncSession, question: str) -> AskResponse:
    ollama_ok = await ollama_service.check_ollama_available()

    # Простая связка с реестром объектов
    if is_registry_objects_question(question):
        objs = await object_service.list_objects(db, overdue_only=True, limit=500)
        if objs:
            lines = [
                f"• {o.name} ({o.municipality}), статус: {o.status.value}, "
                f"плановое окончание: {o.planned_end}"
                for o in objs[:20]
            ]
            answer = (
                f"В реестре найдено просроченных объектов: {len(objs)}.\n\n" + "\n".join(lines)
            )
            if len(objs) > 20:
                answer += f"\n\n… и ещё {len(objs) - 20} (см. фильтр «просроченные» в реестре)."
            return AskResponse(
                answer=answer,
                mode="objects",
                citations=[],
                ollama_available=ollama_ok,
                note="Ответ сформирован запросом к базе объектов, без генерации LLM.",
            )

    hits, search_mode = await search_service.hybrid_search(db, question, top_k=5)
    citations = [
        Citation(
            document_id=ch.document_id,
            document_title=ch.document.title,
            chunk_index=ch.chunk_index,
            excerpt=ch.content[:400] + ("…" if len(ch.content) > 400 else ""),
            score=score,
        )
        for ch, score in hits
    ]

    if not hits:
        return AskResponse(
            answer="В загруженных документах не найдено подходящих фрагментов. "
            "Загрузите нормативные документы или используйте демо-файлы из каталога data/samples.",
            mode=search_mode,
            citations=[],
            ollama_available=ollama_ok,
            note=None if ollama_ok else "Ollama недоступен — используется ключевой поиск.",
        )

    context = "\n\n---\n\n".join(
        f"[Документ: {ch.document.title}, фрагмент {ch.chunk_index + 1}]\n{ch.content}"
        for ch, _ in hits
    )

    if ollama_ok:
        generated = await ollama_service.chat(SYSTEM_PROMPT, question, context)
        if generated:
            return AskResponse(
                answer=generated,
                mode="llm",
                citations=citations,
                ollama_available=True,
            )

    # Fallback без генерации: собираем ответ из найденных фрагментов
    mode_label = "векторный" if search_mode == "vector" else "ключевой"
    parts = [f"Найдены релевантные фрагменты ({mode_label} поиск):\n"]
    for i, (ch, score) in enumerate(hits, 1):
        parts.append(
            f"{i}. «{ch.document.title}», фр. {ch.chunk_index + 1} (релевантность {score:.2f}):\n"
            f"{ch.content[:500]}{'…' if len(ch.content) > 500 else ''}"
        )
    note = (
        "Модель Ollama недоступна или не ответила — показаны фрагменты без генерации текста."
        if not ollama_ok
        else "Генерация не удалась — показаны фрагменты без LLM."
    )
    return AskResponse(
        answer="\n\n".join(parts),
        mode="keyword",
        citations=citations,
        ollama_available=ollama_ok,
        note=note,
    )


async def seed_sample_documents(db: AsyncSession) -> int:
    """Загружает демо-документы из data/samples, если их ещё нет."""
    samples_path = Path(settings.samples_dir)
    if not samples_path.exists():
        return 0
    existing = await db.execute(select(Document).where(Document.is_demo.is_(True)))
    if existing.scalars().first():
        return 0
    count = 0
    for path in sorted(samples_path.glob("*.txt")):
        content = path.read_bytes()
        title = path.stem.replace("_", " ")
        await ingest_bytes(
            db,
            title=f"[ДЕМО] {title}",
            filename=path.name,
            content=content,
            is_demo=True,
        )
        count += 1
    return count
